"""
Helper script to compile text sample resumes into PDF and DOCX files.
"""

from pathlib import Path
from fpdf import FPDF
from docx import Document

RESUMES_DIR = Path(__file__).resolve().parent

samples = [
    "alex_morgan_ml_engineer",
    "sarah_chen_fullstack",
    "david_kumar_data_analyst",
    "incomplete_junior_resume"
]

for name in samples:
    txt_file = RESUMES_DIR / f"{name}.txt"
    if not txt_file.exists():
        continue
    content = txt_file.read_text(encoding="utf-8")

    # 1. Generate DOCX
    docx_file = RESUMES_DIR / f"{name}.docx"
    doc = Document()
    for line in content.splitlines():
        line_str = line.strip()
        if line_str:
            if line_str.isupper() and len(line_str) < 40:
                doc.add_heading(line_str, level=2)
            else:
                doc.add_paragraph(line_str)
    doc.save(docx_file)

    # 2. Generate PDF
    pdf_file = RESUMES_DIR / f"{name}.pdf"
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    epw = pdf.epw

    for line in content.splitlines():
        line_str = line.strip()
        if not line_str:
            pdf.ln(3)
            continue
        clean_line = line_str.encode("latin-1", "replace").decode("latin-1")
        if clean_line.isupper() and len(clean_line) < 40:
            pdf.set_font("Helvetica", style="B", size=12)
            pdf.ln(2)
            pdf.multi_cell(w=epw, h=7, text=clean_line)
            pdf.set_font("Helvetica", size=10)
        else:
            pdf.multi_cell(w=epw, h=5, text=clean_line)

    pdf.output(str(pdf_file))

print(f"Successfully generated PDF and DOCX files in {RESUMES_DIR}")
