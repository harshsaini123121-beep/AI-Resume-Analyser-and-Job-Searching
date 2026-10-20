"""
Document parser supporting PDF, DOCX, and TXT files with fallback mechanisms.
Extracts clean text and document metadata (page count, format, encoding).
"""

import io
import os
from pathlib import Path
import re
from typing import Union, BinaryIO, Dict, Any, Tuple, Optional
import uuid
import zipfile

from app.services.cleaner import clean_text


class ResumeParsingError(Exception):
    """Custom exception raised when resume text extraction fails."""
    pass


def validate_magic_bytes(content: bytes, ext: str) -> None:
    """
    Validates file content headers (magic bytes) to prevent extension-spoofing attacks
    (e.g., Windows executable .exe or binary renamed to .txt or .pdf).
    """
    if len(content) == 0:
        raise ResumeParsingError("The file is empty (0 bytes).")

    ext_lower = ext.lower().strip()
    if not ext_lower.startswith("."):
        ext_lower = f".{ext_lower}"

    if ext_lower == ".pdf":
        if not content.startswith(b"%PDF"):
            raise ResumeParsingError("Invalid PDF header signature. File content is not a valid PDF document.")
    elif ext_lower in (".docx", ".doc"):
        if ext_lower == ".docx":
            if not content.startswith(b"PK\x03\x04"):
                raise ResumeParsingError("Invalid DOCX signature. File content is not a valid OpenXML document.")
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as zf:
                    if "word/document.xml" not in zf.namelist():
                        raise ResumeParsingError("Corrupt DOCX archive: missing word/document.xml structure.")
            except zipfile.BadZipFile:
                raise ResumeParsingError("Invalid or corrupted DOCX archive.")
    elif ext_lower in (".txt", ".md", ".rtf"):
        # Detect binary PE executable header (MZ) or ELF
        if content.startswith(b"MZ") or content.startswith(b"\x7fELF"):
            raise ResumeParsingError("Binary executable content detected. Only plain text resumes are accepted.")
        # Check for binary null bytes in initial sample
        sample = content[:1024]
        if b"\x00" in sample:
            raise ResumeParsingError("Binary content detected in text file (contains NUL bytes).")


def sanitize_filename(filename: str) -> Tuple[str, str]:
    """
    Sanitizes raw uploaded filenames:
    - Generates a UUID for internal file isolation
    - Strips directory traversal components (.., /, \\)
    - Strips HTML tags and non-whitelisted characters
    Returns (storage_uuid, clean_display_name).
    """
    storage_uuid = str(uuid.uuid4())
    raw_name = Path(filename or "resume.pdf").name
    # Strip HTML tags
    clean_name = re.sub(r"<[^>]*>", "", raw_name)
    # Strip invalid path/script characters
    clean_name = re.sub(r"[^\w\.\-\_]", "_", clean_name).strip("._")
    if not clean_name:
        clean_name = "resume.pdf"
    return storage_uuid, clean_name


def _normalize_stream(file_source: Union[str, Path, BinaryIO, bytes]) -> io.BytesIO:
    """Helper to convert any file source into a seekable BytesIO stream."""
    if isinstance(file_source, bytes):
        if len(file_source) == 0:
            raise ResumeParsingError("The file is empty (0 bytes).")
        return io.BytesIO(file_source)
    elif hasattr(file_source, "read"):
        content = file_source.read()
        if len(content) == 0:
            raise ResumeParsingError("The file is empty (0 bytes).")
        stream = io.BytesIO(content)
        if hasattr(file_source, "seek"):
            try:
                file_source.seek(0)
            except Exception:
                pass
        return stream
    elif isinstance(file_source, (str, Path)):
        p = Path(file_source)
        if not p.exists():
            raise ResumeParsingError(f"File not found: {file_source}")
        content = p.read_bytes()
        if len(content) == 0:
            raise ResumeParsingError("The file is empty (0 bytes).")
        return io.BytesIO(content)
    else:
        raise ResumeParsingError(f"Unsupported file source type: {type(file_source)}")


def extract_document_from_pdf(file_source: Union[str, Path, BinaryIO, bytes]) -> Tuple[str, Optional[int]]:
    """
    Extracts text and page count from a PDF file using pdfplumber with PyPDF fallback.
    Handles encrypted, scanned, and malformed documents gracefully.
    """
    stream = _normalize_stream(file_source)
    extracted_pages = []
    page_count: Optional[int] = None

    # 1. Primary parser: pdfplumber
    try:
        import pdfplumber

        stream.seek(0)
        with pdfplumber.open(stream) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                page_text = page.extract_text(layout=True) or ""
                # Also capture text inside tables if layout missed it
                tables = page.extract_tables()
                table_text = ""
                if tables:
                    table_rows = []
                    for table in tables:
                        for row in table:
                            filtered_row = [str(cell).strip() for cell in row if cell]
                            if filtered_row:
                                table_rows.append(" | ".join(filtered_row))
                    table_text = "\n".join(table_rows)

                combined = f"{page_text}\n{table_text}".strip()
                if combined:
                    extracted_pages.append(combined)

    except Exception as e_pdfplumber:
        # 2. Fallback parser: pypdf
        try:
            from pypdf import PdfReader
            from pypdf.errors import FileNotDecryptedError

            stream.seek(0)
            reader = PdfReader(stream)

            if getattr(reader, "is_encrypted", False):
                try:
                    # Attempt decrypt with empty password
                    decrypted = reader.decrypt("")
                    if not decrypted:
                        raise ResumeParsingError("The PDF file is password-protected. Please provide an unencrypted PDF.")
                except Exception:
                    raise ResumeParsingError("The PDF file is password-protected. Please provide an unencrypted PDF.")

            page_count = len(reader.pages)
            extracted_pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    extracted_pages.append(text)
        except ResumeParsingError:
            raise
        except Exception as e_pypdf:
            raise ResumeParsingError(
                f"Failed to extract text from PDF: {e_pdfplumber}; {e_pypdf}"
            )

    full_text = "\n\n".join(extracted_pages).strip()

    if not full_text or len(full_text.strip()) < 20:
        raise ResumeParsingError(
            "The PDF document produced virtually no text. It may be a scanned image or empty. "
            "Please use a text-selectable PDF or DOCX file."
        )

    return full_text, page_count


def extract_text_from_pdf(file_source: Union[str, Path, BinaryIO, bytes]) -> str:
    """Legacy helper returning only the extracted PDF text."""
    text, _ = extract_document_from_pdf(file_source)
    return text


def extract_document_from_docx(file_source: Union[str, Path, BinaryIO, bytes]) -> Tuple[str, Optional[int]]:
    """
    Extracts text from a DOCX file using python-docx.
    Captures both paragraph bodies and tables. Page count is not fixed in DOCX.
    """
    try:
        from docx import Document
    except ImportError:
        raise ResumeParsingError("python-docx is not installed in the environment.")

    stream = _normalize_stream(file_source)

    try:
        stream.seek(0)
        doc = Document(stream)
        elements = []

        # Read paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                elements.append(text)

        # Read tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    unique_cells = []
                    for c in cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    elements.append(" | ".join(unique_cells))

        full_text = "\n".join(elements).strip()
        if not full_text or len(full_text.strip()) < 10:
            raise ResumeParsingError("The DOCX file is empty or contains no readable text.")

        return full_text, None

    except ResumeParsingError:
        raise
    except Exception as e:
        raise ResumeParsingError(f"Error reading DOCX document: {str(e)}")


def extract_text_from_docx(file_source: Union[str, Path, BinaryIO, bytes]) -> str:
    """Legacy helper returning only the extracted DOCX text."""
    text, _ = extract_document_from_docx(file_source)
    return text


def extract_document_from_txt(file_source: Union[str, Path, BinaryIO, bytes]) -> Tuple[str, Optional[int]]:
    """
    Extracts text from a plain text file, handling diverse encodings.
    """
    stream = _normalize_stream(file_source)
    raw_bytes = stream.getvalue()

    if not raw_bytes or len(raw_bytes.strip()) == 0:
        raise ResumeParsingError("The text file is empty (0 bytes).")

    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            decoded = raw_bytes.decode(encoding).strip()
            if not decoded or len(decoded) < 10:
                raise ResumeParsingError("The text file contains no readable content.")
            return decoded, None
        except UnicodeDecodeError:
            continue

    raise ResumeParsingError("Failed to decode text file with standard encodings.")


def extract_text_from_txt(file_source: Union[str, Path, BinaryIO, bytes]) -> str:
    """Legacy helper returning only the extracted TXT text."""
    text, _ = extract_document_from_txt(file_source)
    return text


def extract_document(
    file_source: Union[str, Path, BinaryIO, bytes],
    filename: str
) -> Dict[str, Any]:
    """
    Unified router that selects the proper parser based on the file extension,
    validates magic bytes, extracts the raw text, validates non-empty content,
    and returns text with sanitized metadata.
    """
    if not filename:
        raise ResumeParsingError("Filename cannot be empty.")

    storage_uuid, clean_filename = sanitize_filename(filename)
    ext = os.path.splitext(clean_filename)[1].lower()

    # Normalize stream and read raw bytes for magic-byte validation
    stream = _normalize_stream(file_source)
    content_bytes = stream.getvalue()
    validate_magic_bytes(content_bytes, ext)

    if ext == ".pdf":
        raw_text, page_count = extract_document_from_pdf(content_bytes)
    elif ext in (".docx", ".doc"):
        if ext == ".doc":
            raise ResumeParsingError(
                "Legacy .doc binary format is not supported. Please save/convert as .docx or .pdf."
            )
        raw_text, page_count = extract_document_from_docx(content_bytes)
    elif ext in (".txt", ".md", ".rtf"):
        raw_text, page_count = extract_document_from_txt(content_bytes)
    else:
        raise ResumeParsingError(
            f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .txt"
        )

    cleaned = clean_text(raw_text)
    if not cleaned or len(cleaned.strip()) < 10:
        raise ResumeParsingError("The document produced no readable text. Please check the document.")

    return {
        "text": cleaned,
        "page_count": page_count,
        "format": ext,
        "raw_length": len(raw_text),
        "storage_uuid": storage_uuid,
        "filename": clean_filename
    }


def extract_text(file_source: Union[str, Path, BinaryIO, bytes], filename: str) -> str:
    """
    Backward-compatible extraction function returning cleaned text string.
    """
    doc = extract_document(file_source, filename)
    return doc["text"]
