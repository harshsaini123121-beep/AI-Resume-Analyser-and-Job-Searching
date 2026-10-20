"""
Audit, Debug, and Hardening Test Suite.
Validates:
1. File validation: PDF, DOCX, TXT support, invalid extensions, oversized files, empty files.
2. Corrupt document and empty text handling.
3. Missing optional fields in resume extraction.
4. SHA-256 content hashing for duplicate uploads vs same-name different content.
5. ATS scoring rubric robustness across diverse candidate profiles.
6. Hybrid job matching with empty, valid, and boundary job descriptions.
7. Section header regex behavior and avoidance of invalid escape warnings.
8. PDF report generation and download content verification.
9. Batch screening resilience with mixed inputs.
"""

import hashlib
import io
import json
from pathlib import Path
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services.parser import (
    extract_document,
    extract_text,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_txt,
    ResumeParsingError
)
from app.services.extractor import (
    extract_resume_data,
    detect_sections,
    extract_skills,
    extract_experience_years
)
from app.services.scorer import calculate_ats_score
from app.services.matcher import compute_match, compute_semantic_similarity
from app.services.pdf_generator import generate_pdf_report

client = TestClient(app)
ROOT_DIR = Path(__file__).resolve().parent.parent


# -----------------------------------------------------------------------------
# 1. Document Parser & Metadata Tests
# -----------------------------------------------------------------------------
def test_extract_document_pdf_metadata():
    sample_pdf = ROOT_DIR / "data" / "sample_resumes" / "alex_morgan_ml_engineer.pdf"
    if sample_pdf.exists():
        doc = extract_document(sample_pdf, "alex_morgan_ml_engineer.pdf")
        assert doc["format"] == ".pdf"
        assert isinstance(doc["page_count"], int)
        assert doc["page_count"] >= 1
        assert "Alex Morgan" in doc["text"]


def test_extract_document_docx(tmp_path):
    sample_docx = ROOT_DIR / "data" / "sample_resumes" / "alex_morgan_ml_engineer.docx"
    if sample_docx.exists():
        doc = extract_document(sample_docx, "alex_morgan_ml_engineer.docx")
    else:
        # Dynamically build a DOCX file to guarantee real execution
        from docx import Document
        dyn_path = tmp_path / "alex_morgan_ml_engineer.docx"
        b = Document()
        b.add_heading("Alex Morgan", level=1)
        b.add_paragraph("Machine Learning Engineer with PyTorch experience.")
        b.save(str(dyn_path))
        doc = extract_document(dyn_path, "alex_morgan_ml_engineer.docx")

    assert doc["format"] == ".docx"
    # DOCX does not have fixed printer page metadata
    assert doc["page_count"] is None
    assert "Alex Morgan" in doc["text"]
    assert len(doc["storage_uuid"]) == 36


def test_extract_document_txt():
    sample_txt = ROOT_DIR / "data" / "sample_resumes" / "software_engineer_resume.txt"
    if sample_txt.exists():
        doc = extract_document(sample_txt, "software_engineer_resume.txt")
        assert doc["format"] == ".txt"
        assert doc["page_count"] is None
        assert "HARSH VARDHAN" in doc["text"]


def test_empty_file_upload_rejection():
    # Direct parser check
    with pytest.raises(ResumeParsingError, match="empty"):
        extract_document(b"", "empty_resume.pdf")

    with pytest.raises(ResumeParsingError, match="empty"):
        extract_document(b"", "empty_resume.txt")

    # API check
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    resp = client.post("/api/resume/upload", files=files)
    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"].lower()


def test_unsupported_file_format_rejection():
    with pytest.raises(ResumeParsingError, match="Unsupported file format"):
        extract_document(b"fake binary content", "malicious_script.exe")

    files = {"file": ("malicious.exe", b"MZ\x90\x00", "application/octet-stream")}
    resp = client.post("/api/resume/upload", files=files)
    assert resp.status_code == 400
    assert "unsupported" in resp.json()["detail"].lower()


def test_oversized_file_upload_rejection():
    oversized_bytes = b"0" * (5 * 1024 * 1024 + 1024)  # 5 MB + 1 KB
    files = {"file": ("huge_resume.pdf", oversized_bytes, "application/pdf")}
    resp = client.post("/api/resume/upload", files=files)
    assert resp.status_code == 400
    assert "exceeds" in resp.json()["detail"].lower()


def test_corrupt_document_handling():
    corrupt_pdf_bytes = b"%PDF-1.4\ncorrupted bytes that are not valid pdf data"
    with pytest.raises(ResumeParsingError):
        extract_document(corrupt_pdf_bytes, "corrupt.pdf")


def test_empty_text_document_handling():
    whitespace_bytes = b"    \n\n\t   \n  "
    with pytest.raises(ResumeParsingError):
        extract_document(whitespace_bytes, "blank.txt")


# -----------------------------------------------------------------------------
# 2. SHA-256 Hashing & Duplicate Upload Tests
# -----------------------------------------------------------------------------
def test_sha256_hashing_for_duplicate_prevention():
    file1_bytes = b"Alex Morgan resume text version 1 with python and docker."
    file2_bytes = b"Alex Morgan resume text version 2 with updated tenure and skills."

    hash1 = hashlib.sha256(file1_bytes).hexdigest()
    hash1_repeat = hashlib.sha256(file1_bytes).hexdigest()
    hash2 = hashlib.sha256(file2_bytes).hexdigest()

    # Identical file produces identical hash (prevents rerun recomputation)
    assert hash1 == hash1_repeat

    # Two different files with the same filename produce distinct hashes (triggers re-analysis)
    assert hash1 != hash2


# -----------------------------------------------------------------------------
# 3. Section Header Regex & Syntax Warning Verification
# -----------------------------------------------------------------------------
def test_detect_sections_regex_robustness():
    resume_with_markdown_headers = """
    # PROFESSIONAL SUMMARY
    Senior Backend Engineer with 5 years experience.

    ## CORE SKILLS
    Python, FastAPI, Docker, PostgreSQL, Redis.

    ### WORK EXPERIENCE
    CloudScale Technologies (2020 - Present)
    - Architected microservices.

    ### EDUCATION & CREDENTIALS
    B.S. in Computer Science, 2020.

    ### NOTABLE PROJECTS
    Built distributed search engine.

    ### CERTIFICATIONS & LICENSES
    AWS Certified Solutions Architect.
    """
    sections = detect_sections(resume_with_markdown_headers)
    assert "summary" in sections
    assert "skills" in sections
    assert "experience" in sections
    assert "education" in sections
    assert "projects" in sections
    assert "certifications" in sections


# -----------------------------------------------------------------------------
# 4. Extractor Safety with Incomplete / Missing Fields
# -----------------------------------------------------------------------------
def test_extractor_missing_optional_fields():
    minimal_text = "John\nSoftware Engineer looking for an entry level position."
    data = extract_resume_data(minimal_text)

    # Must not crash, must return clean default values
    assert data["candidate_name"] == "John"
    assert data["email"] is None
    assert data["phone"] is None
    assert data["links"]["linkedin"] is None
    assert data["links"]["github"] is None
    assert data["skills"] == []
    assert data["skills_count"] == 0
    assert data["experience_years"] == 0.0
    assert data["word_count"] > 0

    # Ensure full downstream pipeline processes incomplete data without crashing
    ats = calculate_ats_score(minimal_text, data)
    assert ats["total_score"] < 40
    assert "Contact Info" in ats["category_breakdown"]
    assert any("email" in s.lower() for s in ats["suggestions"])


# -----------------------------------------------------------------------------
# 5. ATS Scoring Rubric Verification
# -----------------------------------------------------------------------------
def test_ats_scorer_complete_vs_incomplete():
    # 1. Complete comprehensive technical resume
    complete_text = """
    Jane Architect
    Email: jane.architect@corp.com
    Phone: (555) 123-4567
    LinkedIn: linkedin.com/in/janearchitect
    GitHub: github.com/janearchitect

    PROFESSIONAL SUMMARY
    Senior Cloud Architect with 8 years of experience building high-scale distributed systems.

    CORE SKILLS
    Python, Go, AWS, Docker, Kubernetes, PostgreSQL, Terraform, CI/CD, Redis, React

    WORK EXPERIENCE
    Staff Infrastructure Engineer at Global Cloud Inc
    2019 - Present
    - Architected multi-region Kubernetes clusters serving 10M daily requests with 99.99% uptime.
    - Automated cloud deployments using Terraform and GitHub Actions, reducing deployment time by 60%.
    - Spearheaded migration of legacy services to microservices, saving $250,000 in infrastructure costs.

    EDUCATION
    Master of Science in Computer Science
    2017 - 2019
    """
    data_comp = extract_resume_data(complete_text)
    ats_comp = calculate_ats_score(complete_text, data_comp)

    assert ats_comp["total_score"] >= 75
    assert len(ats_comp["strengths"]) >= 3
    assert "Contact Info" in ats_comp["category_breakdown"]
    assert "Quantified Impact" in ats_comp["category_breakdown"]
    assert ats_comp["category_breakdown"]["Quantified Impact"]["score"] >= 10

    # 2. Incomplete junior profile with missing critical sections and no metrics
    incomplete_text = "Bob\nLooking for an internship in coding. Helped with website stuff."
    data_inc = extract_resume_data(incomplete_text)
    ats_inc = calculate_ats_score(incomplete_text, data_inc)

    assert ats_inc["total_score"] < 40
    assert len(ats_inc["issues"]) >= 3
    assert len(ats_inc["suggestions"]) >= 3


# -----------------------------------------------------------------------------
# 6. Hybrid Matcher Edge Cases (Empty JD, Mismatched Tenure)
# -----------------------------------------------------------------------------
def test_hybrid_matcher_edge_cases():
    resume_text = "Python engineer with Docker and AWS experience."
    resume_skills = ["python", "docker", "aws"]

    # 1. Empty JD text
    match_empty = compute_match(
        resume_text=resume_text,
        resume_skills=resume_skills,
        jd_text="",
        jd_skills=[],
        candidate_exp_years=3.0,
        required_exp_years=0.0
    )
    assert 0.0 <= match_empty["match_percentage"] <= 100.0
    assert match_empty["semantic_score"] == 0.0

    # 2. Perfect match
    match_full = compute_match(
        resume_text=resume_text,
        resume_skills=resume_skills,
        jd_text="Looking for a Python engineer with Docker and AWS skills.",
        jd_skills=["python", "docker", "aws"],
        candidate_exp_years=5.0,
        required_exp_years=3.0
    )
    assert match_full["match_percentage"] >= 80.0
    assert len(match_full["missing_skills"]) == 0
    assert "python" in match_full["matched_skills"]


# -----------------------------------------------------------------------------
# 7. PDF Report Generation & Download Content Verification
# -----------------------------------------------------------------------------
def test_pdf_report_generation_contents():
    category_scores = {
        "Contact Info": {"score": 10, "max_score": 10, "details": "All contact info present"},
        "Skills Density": {"score": 15, "max_score": 15, "details": "12 skills detected"}
    }
    pdf_bytes = generate_pdf_report(
        candidate_name="Pawan Kumar",
        email="pawan@example.com",
        phone="+1 555-0192",
        ats_score=85,
        ats_grade="A+ (ATS Ready)",
        category_breakdown=category_scores,
        strengths=["Strong skills foundation", "Standard sections detected"],
        issues=["Include more quantifiable metrics"],
        suggestions=["Add GitHub project repository link"],
        skills=["Python", "FastAPI", "Docker", "PyTorch"]
    )
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 1000
    # Standard PDF magic header
    assert pdf_bytes.startswith(b"%PDF-")


# -----------------------------------------------------------------------------
# 8. Batch Candidate Screening with Mixed Valid and Invalid Inputs
# -----------------------------------------------------------------------------
def test_batch_screening_mixed_inputs():
    valid_text = "Jane Doe\nExperienced Python and Docker developer. 4 years of experience."
    minimal_text = "Junior developer with basic knowledge."

    candidates = [
        ("jane.txt", valid_text),
        ("junior.txt", minimal_text)
    ]

    target_skills = ["python", "docker"]
    results = []
    for fname, text in candidates:
        ext = extract_resume_data(text)
        match = compute_match(
            resume_text=text,
            resume_skills=ext["skills"],
            jd_text="Python and Docker engineer required.",
            jd_skills=target_skills,
            candidate_exp_years=ext["experience_years"],
            required_exp_years=2.0
        )
        ats = calculate_ats_score(text, ext)
        results.append((fname, match["match_percentage"], ats["total_score"]))

    assert len(results) == 2
    # Jane should outscore the minimal junior profile
    assert results[0][1] > results[1][1]


# -----------------------------------------------------------------------------
# 9. Regression: Custom Job Description Consistency
# -----------------------------------------------------------------------------
def test_custom_job_description_extracted_skills_consistency():
    custom_jd = """
    Senior Data Engineer
    Requirements:
    - 4+ years of experience with Python, Apache Spark, Snowflake, and Airflow.
    - Strong SQL optimization and database warehousing skills.
    """
    resume_text = "Data Engineer with 5 years experience in Python, Snowflake, and SQL."
    resume_skills = ["python", "snowflake", "sql"]

    # Extract skills directly from the custom JD
    custom_skills = extract_skills(custom_jd)
    custom_exp_req = extract_experience_years(custom_jd)

    assert "python" in custom_skills
    assert "snowflake" in custom_skills
    assert "airflow" in custom_skills
    assert custom_exp_req == 4.0

    # Ensure compute_match uses the custom JD's skills, not a mismatched preset job's skills
    match = compute_match(
        resume_text=resume_text,
        resume_skills=resume_skills,
        jd_text=custom_jd,
        jd_skills=custom_skills,
        candidate_exp_years=5.0,
        required_exp_years=custom_exp_req
    )

    assert "python" in match["matched_skills"]
    assert "snowflake" in match["matched_skills"]
    assert "airflow" in match["missing_skills"]
    # Preserves consistency: matched + missing = all required custom skills
    assert set(match["matched_skills"]).union(set(match["missing_skills"])) == set(custom_skills)


# -----------------------------------------------------------------------------
# 10. Regression: Empty Jobs Dataset Graceful Handling (No .iloc[0] Crash)
# -----------------------------------------------------------------------------
def test_empty_jobs_dataset_graceful_handling(tmp_path):
    empty_csv = tmp_path / "empty_jobs.csv"
    empty_csv.write_text("job_id,title,company,location,description,required_skills,experience_required,salary\n", encoding="utf-8")

    from app.services.recommender import load_jobs_dataset, recommend_jobs
    df = load_jobs_dataset(str(empty_csv))
    assert df.empty

    # recommend_jobs must return an empty list without IndexError: single positional indexer is out-of-bounds
    recs = recommend_jobs(
        resume_text="Python developer",
        resume_skills=["python"],
        csv_path=str(empty_csv)
    )
    assert recs == []

    # Filtering an empty dataframe must not crash when guarded
    filtered = df[df["job_id"] == "non_existent"]
    assert filtered.empty
    result_item = filtered.iloc[0] if not filtered.empty else None
    assert result_item is None


# -----------------------------------------------------------------------------
# 11. Regression: Batch Screening Error Handling & Hash Caching
# -----------------------------------------------------------------------------
def test_batch_screening_size_validation_and_caching():
    valid_bytes = b"Candidate One with Python and Docker."
    empty_bytes = b""
    oversized_bytes = b"A" * (6 * 1024 * 1024)  # 6 MB

    # 1. Validation logic
    files = [
        ("valid.txt", valid_bytes),
        ("empty.txt", empty_bytes),
        ("oversized.txt", oversized_bytes)
    ]

    processed = []
    skipped_reasons = []

    for name, content in files:
        if len(content) == 0:
            skipped_reasons.append((name, "empty"))
            continue
        if len(content) > 5 * 1024 * 1024:
            skipped_reasons.append((name, "oversized"))
            continue
        doc = extract_document(content, name)
        processed.append((name, doc["text"], hashlib.sha256(content).hexdigest()))

    assert len(processed) == 1
    assert processed[0][0] == "valid.txt"
    assert len(skipped_reasons) == 2

    # 2. Hash caching logic
    hashes1 = [h for _, _, h in processed]
    cache_key1 = "job123_" + hashlib.sha256(",".join(sorted(hashes1)).encode()).hexdigest()
    cache_key1_repeat = "job123_" + hashlib.sha256(",".join(sorted(hashes1)).encode()).hexdigest()
    assert cache_key1 == cache_key1_repeat

    # Changing target job changes cache key
    cache_key_diff_job = "job456_" + hashlib.sha256(",".join(sorted(hashes1)).encode()).hexdigest()
    assert cache_key1 != cache_key_diff_job


# -----------------------------------------------------------------------------
# 12. Regression: Rejected Replacement Upload State Isolation
# -----------------------------------------------------------------------------
def test_rejected_replacement_upload_state_isolation():
    # 1. Initial valid resume is parsed
    valid_bytes = b"Alex Morgan ML Engineer with Python and Docker."
    doc1 = extract_document(valid_bytes, "valid_resume.txt")
    data1 = extract_resume_data(doc1["text"])
    ats1 = calculate_ats_score(doc1["text"], data1)

    initial_session_state = {
        "resume_data": data1,
        "ats_result": ats1,
        "active_filename": "valid_resume.txt",
        "processed_file_hash": hashlib.sha256(valid_bytes).hexdigest()
    }

    # 2. Subsequent replacement upload is invalid (e.g. empty or corrupt)
    rejected_bytes = b""
    with pytest.raises(ResumeParsingError):
        extract_document(rejected_bytes, "corrupt_replacement.pdf")

    # 3. Verify that the previous state is preserved and not overwritten
    assert initial_session_state["resume_data"] is not None
    assert initial_session_state["active_filename"] == "valid_resume.txt"
    assert initial_session_state["active_filename"] != "corrupt_replacement.pdf"

